from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, PageBreak
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT = Path(__file__).resolve().parents[2]

def rows(meeting):
    yield ['№', 'Поручение', 'Ответственный', 'Срок']
    for i,a in enumerate(meeting['actions'],1):
        yield [str(i), a['title'], a.get('owner') or 'Не определён', a.get('due_date') or a.get('deadline_text') or 'Не указан']

def export_docx(m):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.75)
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10)
    doc.core_properties.author = 'GestSpeak'
    doc.add_heading('Протокол совещания', 0)
    doc.add_paragraph(m['title'])
    doc.add_paragraph(f"Дата: {m['meeting_date']} · {m.get('organization','')}")
    doc.add_paragraph('Статус: ' + ('ЧЕРНОВИК — требуется проверка' if any(a['needs_review'] for a in m['actions']) else 'Поручения проверены секретарём'))
    if m['source'] == 'demo':
        doc.add_paragraph('Демонстрационный размеченный пример. Не является результатом распознавания аудио.')
    doc.add_heading('Краткое содержание',1)
    for p in m['summary']:
        doc.add_paragraph(p, style='List Bullet')
    doc.add_heading('Поручения',1)
    table = doc.add_table(rows=0, cols=4)
    table.autofit = False
    table.style = 'Light Shading Accent 1'
    for row in rows(m):
        cells = table.add_row().cells
        for cell,text,width in zip(cells,row,[.35,3.1,2,1.55]):
            cell.text = text
            cell.width = Inches(width)
    repeat = OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_heading('Источники поручений',1)
    for i,a in enumerate(m['actions'],1):
        doc.add_paragraph(f"{i}. {', '.join(a['evidence'])}: {a['quote']}")
        if a['needs_review']: doc.add_paragraph('Проверить: ' + a['review_reason'])
    doc.add_page_break()
    doc.add_heading('Транскрипт',1)
    for s in m['segments']:
        p = doc.add_paragraph()
        p.add_run(f"[{s['id']}] {s['speaker']}: ").bold = True
        p.add_run(s['text'])
    stream = BytesIO(); doc.save(stream); return stream.getvalue()

def export_pdf(m):
    if 'NotoSans' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('NotoSans', str(ROOT/'assets/fonts/NotoSans-Regular.ttf')))
    style = ParagraphStyle('body', fontName='NotoSans', fontSize=9, leading=14, spaceAfter=7)
    title = ParagraphStyle('title', parent=style, fontSize=23, leading=29, spaceAfter=14)
    heading = ParagraphStyle('heading', parent=style, fontSize=14, leading=20, spaceBefore=15, spaceAfter=10)
    def p(text, st=style): return Paragraph(escape(str(text)), st)
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=40, leftMargin=40, topMargin=42, bottomMargin=42)
    content = [p('Протокол совещания', title), p(m['title']), p(m['meeting_date']+' · '+m.get('organization',''))]
    if m['source']=='demo': content.append(p('Демонстрационный размеченный пример. Не результат распознавания аудио.'))
    content.append(p('ЧЕРНОВИК — требуется проверка' if any(a['needs_review'] for a in m['actions']) else 'Поручения проверены секретарём'))
    content.append(p('Краткое содержание',heading))
    content += [p(text) for text in m['summary']]
    content.append(p('Поручения',heading))
    table = LongTable([[p(v) for v in row] for row in rows(m)], colWidths=[24,239,140,112-0.72], repeatRows=1, hAlign='LEFT')
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eeeeee')),('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,-1),.4,colors.HexColor('#dddddd')),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8)]))
    content.append(table)
    content.append(p('Источники поручений',heading))
    for i,a in enumerate(m['actions'],1):
        content.append(p(f"{i}. {', '.join(a['evidence'])}: {a['quote']}"))
        if a['needs_review']: content.append(p('Проверить: '+a['review_reason']))
    content += [PageBreak(),p('Транскрипт',heading)]
    content += [p(f"[{s['id']}] {s['speaker']}: {s['text']}") for s in m['segments']]
    def footer(canvas, doc):
        canvas.setFont('NotoSans',8); canvas.drawString(40,22,'GestSpeak · Внутренний документ'); canvas.drawRightString(A4[0]-40,22,str(doc.page))
    doc.build(content,onFirstPage=footer,onLaterPages=footer)
    return output.getvalue()
