import os
import tempfile
os.environ['GESTSPEAK_DATA']=tempfile.mkdtemp(prefix='gestspeak-tests-')
from datetime import date
from io import BytesIO
import pytest
from fastapi.testclient import TestClient
from gestspeak.dates import resolve_deadline
from gestspeak.engine import attach_speakers, parse_transcript, validate_analysis, local_url, rule_extract
from gestspeak.schema import Analysis, Action, Segment
from gestspeak.demo import make_demo
from gestspeak.main import app
from gestspeak.exports import export_pdf, export_docx
from pypdf import PdfReader
from docx import Document

@pytest.mark.parametrize('raw,expected', [('до 15 октября','2026-10-15'),('К пятнадцатому октября','2026-10-15'),('жұма','2026-09-25'),('к среде','2026-09-23'),('30 қыркүйек','2026-09-30'),('ертең','2026-09-24'),('на следующей неделе','2026-10-02')])
def test_dates(raw,expected):
    assert str(resolve_deadline(raw,date(2026,9,23))[0])==expected

def test_unknown_and_invalid_dates():
    assert resolve_deadline('позже',date(2026,9,23))[0] is None
    assert resolve_deadline('31 февраля',date(2026,9,23))[0] is None
    assert resolve_deadline('',date(2026,9,23))[0] is None

def test_week_duration_and_kazakh_sunday():
    assert resolve_deadline('жексенбі',date(2026,9,23))[0] == date(2026,9,27)
    assert resolve_deadline('через две недели',date(2026,9,23))[0] == date(2026,10,7)
    assert resolve_deadline('на неделе',date(2026,9,23))[0] is None

def test_never_invent_owner_or_due_date():
    segments=[Segment(id='s1',speaker='A',text='Подготовить отчёт.')]
    result=validate_analysis(Analysis(actions=[Action(title='Подготовить отчёт',owner='Несуществующий',due_date=date(2030,1,1),deadline_text='завтра',evidence=['s1'],quote='выдумка')]),segments,date(2026,9,23),[])
    assert result.actions[0].owner is None
    assert result.actions[0].due_date is None
    assert result.actions[0].needs_review

def test_speaker_boundary_and_silence():
    result=attach_speakers([{'start':0,'end':1,'word':'Привет'}, {'start':1,'end':2,'word':' всем'}, {'start':2,'end':3,'word':' Сәлем'}, {'start':8,'end':9,'word':' Да'}],[(0,2,'SPEAKER_00'),(2,4,'SPEAKER_01')])
    assert [s.speaker for s in result]==['SPEAKER_00','SPEAKER_01','SPEAKER_UNKNOWN']
    assert result[0].text=='Привет всем'

def test_external_api_blocked(monkeypatch):
    for url in ['https://api.openai.com','http://8.8.8.8:11434','http://evil.example','http://127.0.0.1@evil.example']:
        monkeypatch.setenv('OLLAMA_URL',url)
        with pytest.raises(ValueError): local_url()

def test_ru_kk_mixed_explicit_extraction():
    for text, owner in [('А: Проверить контракт — ответственный Айнур, срок до 30 сентября.','Айнур'), ('А: Есеп дайындау — жауапты Нұрлан, мерзімі 15 қазан.','Нұрлан'),('А: Подготовить есеп — ответственный Нурлан, срок ертең.','Нурлан')]:
        result=rule_extract(parse_transcript(text))
        assert len(result.actions)==1
        assert result.actions[0].owner==owner

def test_demo_has_ten_and_final_deadline():
    demo=make_demo()
    assert len(demo['actions'])==10
    assert demo['actions'][6]['due_date']=='2026-10-15'
    assert 's22' in demo['actions'][6]['evidence']
    assert demo['source']=='demo'
    assert all(a['needs_review'] for a in demo['actions'])

def test_export_real_formats_and_kazakh():
    m=make_demo()
    m['summary'].append('Қазақ тілі: Ә Ғ Қ Ң Ө Ұ Ү Һ І')
    pdf=export_pdf(m)
    text=''.join(p.extract_text() for p in PdfReader(BytesIO(pdf)).pages)
    assert 'Қазақ тілі' in text and 'Гульмира' in text and '2026-10-15' in text
    doc=Document(BytesIO(export_docx(m)))
    assert len(doc.tables[0].rows)==11
    assert 'Қазақ тілі' in '\n'.join(p.text for p in doc.paragraphs)

def test_api_review_persistence_export_and_origin():
    with TestClient(app) as c:
        m=c.post('/api/demo').json()
        assert len(c.get('/api/meetings').json())>=1
        patch=c.patch('/api/meetings/demo-samruk/actions/a1',json={'status':'done','needs_review':False})
        assert patch.status_code==200
        assert c.get('/api/meetings/demo-samruk').json()['actions'][0]['status']=='done'
        assert c.get('/api/meetings/demo-samruk/export/pdf').content.startswith(b'%PDF')
        assert c.get('/api/meetings/demo-samruk/export/docx').content.startswith(b'PK')
        assert c.get('/api/meetings/missing').status_code==404
        assert c.post('/api/demo',headers={'Origin':'https://attacker.example'}).status_code==403
        assert c.post('/api/meetings/text',json={'title':'Тест','meeting_date':'2026-09-23','transcript':'Нурлан: Подготовить отчёт до пятницы.'}).status_code==400
        assert c.get('/api/meetings/demo-samruk/audit').json()

def test_auth(monkeypatch):
    monkeypatch.setenv('GESTSPEAK_API_TOKEN','test-secret')
    with TestClient(app) as c:
        assert c.get('/api/meetings').status_code==401
        assert c.get('/api/meetings',headers={'Authorization':'Bearer test-secret'}).status_code==200
